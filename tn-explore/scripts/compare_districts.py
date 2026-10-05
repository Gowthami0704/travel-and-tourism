import openpyxl
import sqlite3
from collections import defaultdict

conn = sqlite3.connect('database/database.sqlite')
cursor = conn.cursor()
cursor.execute("SELECT id, name FROM districts ORDER BY id")
db_districts = cursor.fetchall()

wb = openpyxl.load_workbook('data/Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx', data_only=True)
ws = wb['Master 2500' if 'Master 2500' in wb.sheetnames else wb.sheetnames[0]]
headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
h_map = {h: i for i, h in enumerate(headers)}

excel_districts = defaultdict(list)
for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(row): continue
    d = str(row[h_map.get('District', 0)] or '').strip()
    name = str(row[h_map.get('Name', 1)] or '').strip()
    rtype = str(row[h_map.get('Record Type', 2)] or '').strip()
    if 'place' in rtype.lower():
        excel_districts[d].append(name)

print("DB Districts vs Excel Districts:")
excel_names = list(excel_districts.keys())

for db_id, db_name in db_districts:
    # normalize and find match
    match = None
    db_clean = db_name.lower().replace(' ', '').replace('the', '').replace('-', '')
    for ex in excel_names:
        ex_clean = ex.lower().replace(' ', '').replace('the', '').replace('-', '')
        if ex_clean == db_clean or (db_clean in ex_clean) or (ex_clean in db_clean):
            match = ex
            break
    print(f"DB [{db_id:2d}] {db_name:20s} -> Excel: {str(match):20s} (Count: {len(excel_districts.get(match, []))})")
