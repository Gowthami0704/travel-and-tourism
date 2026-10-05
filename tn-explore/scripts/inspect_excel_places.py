import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('data/Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx', data_only=True)
ws = wb['Master 2500' if 'Master 2500' in wb.sheetnames else wb.sheetnames[0]]
headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
h_map = {h: i for i, h in enumerate(headers)}

dist_places = defaultdict(list)
for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(row): continue
    d = str(row[h_map.get('District', 0)] or '').strip()
    name = str(row[h_map.get('Name', 1)] or '').strip()
    rtype = str(row[h_map.get('Record Type', 2)] or '').strip()
    ptype = str(row[h_map.get('Place/Food/Hotel Type', 3)] or '').strip()
    cat = str(row[h_map.get('Category', 4)] or '').strip()
    desc = str(row[h_map.get('Description', 5)] or '').strip()
    img = str(row[h_map.get('Image URL', 8)] or '').strip()
    lat = row[h_map.get('Latitude', 9)] if 'Latitude' in h_map else None
    lng = row[h_map.get('Longitude', 10)] if 'Longitude' in h_map else None
    
    if 'place' in rtype.lower():
        dist_places[d].append({
            'name': name,
            'ptype': ptype,
            'cat': cat,
            'desc': desc,
            'img': img,
            'lat': lat,
            'lng': lng
        })

print(f"Total districts with places: {len(dist_places)}")
for d in list(dist_places.keys())[:5]:
    print(f"\n--- {d} ({len(dist_places[d])} places) ---")
    for p in dist_places[d][:10]:
        print(f"  • {p['name']} [{p['ptype']} / {p['cat']}] - Img: {p['img'][:40] if p['img'] else 'EMPTY'}")
