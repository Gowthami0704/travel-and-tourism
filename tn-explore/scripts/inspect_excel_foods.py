import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('data/Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx', data_only=True)
ws = wb['Master 2500' if 'Master 2500' in wb.sheetnames else wb.sheetnames[0]]
headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
h_map = {h: i for i, h in enumerate(headers)}

food_items = defaultdict(list)
for row in ws.iter_rows(min_row=2, values_only=True):
    if not any(row): continue
    d = str(row[h_map.get('District')] or '').strip()
    r = str(row[h_map.get('Record Type')] or '').strip()
    name = str(row[h_map.get('Name')] or '').strip()
    shop = str(row[h_map.get('Food Shop / Restaurant')] or '').strip()
    spec = str(row[h_map.get('Food Specialty')] or '').strip()
    desc = str(row[h_map.get('Description')] or '').strip()
    if 'food' in r.lower():
        food_items[d].append({'name': name, 'shop': shop, 'spec': spec, 'desc': desc})

print(f"Districts with food rows: {len(food_items)}")
for d in list(food_items.keys())[:8]:
    print(f"\n=== {d} ({len(food_items[d])} food rows) ===")
    for f in food_items[d][:5]:
        print(f"  • Dish: {f['name']} | Shop: {f['shop']} | Spec: {f['spec']}")
