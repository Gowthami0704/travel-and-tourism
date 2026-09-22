"""
Script 1: Convert Excel Dataset to JSON
Reads the 'Master 2500' sheet from Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx
and generates tourism_data.json with mapped keys.
"""

import os
import json
import openpyxl

EXCEL_PATHS = [
    os.path.join(os.getcwd(), 'data', 'Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx'),
    os.path.join(os.path.dirname(os.getcwd()), 'data', 'Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx'),
    r'C:\Users\princ\Downloads\Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx'
]

OUTPUT_PATHS = [
    os.path.join(os.getcwd(), 'public', 'data', 'tourism_data.json'),
    os.path.join(os.path.dirname(os.getcwd()), 'public', 'data', 'tourism_data.json'),
]

def find_excel_file():
    for p in EXCEL_PATHS:
        if os.path.exists(p):
            return p
    raise FileNotFoundError("Could not locate Tamil_Nadu_Smart_Tourism_2500_Filled_Dataset.xlsx in data/ or Downloads.")

def convert_excel_to_json():
    excel_file = find_excel_file()
    print(f"Reading dataset from: {excel_file}")

    wb = openpyxl.load_workbook(excel_file, data_only=True)
    sheet_name = 'Master 2500' if 'Master 2500' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    print(f"Reading sheet: {sheet_name}")

    headers = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[1]]
    
    # Map headers to indices
    header_indices = {h: idx for idx, h in enumerate(headers)}

    def get_val(row, col_name, default=''):
        idx = header_indices.get(col_name)
        if idx is not None and idx < len(row):
            val = row[idx]
            return str(val).strip() if val is not None else default
        return default

    json_records = []
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not any(row):
            continue

        rec_id = get_val(row, 'Record ID', f"TN-{row_idx:04d}")
        district = get_val(row, 'District', '')
        name = get_val(row, 'Name', '')
        record_type = get_val(row, 'Record Type', 'Place')
        type_val = get_val(row, 'Place/Food/Hotel Type', '')
        category = get_val(row, 'Category', '')
        description = get_val(row, 'Description', '')
        tags = get_val(row, 'AI Recommendation Tags', '')
        maps_ref = get_val(row, 'Google Maps / Location Reference', '')

        # Construct mapped object adhering to specification
        record = {
            "id": rec_id,
            "district": district,
            "name": name,
            "record_type": record_type,
            "type": type_val,
            "category": category,
            "description": description,
            "tags": tags,
            "maps_url": maps_ref,
            "image_url": ""
        }
        json_records.append(record)

    print(f"Successfully converted {len(json_records)} records.")

    for out_path in OUTPUT_PATHS:
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        with open(out_path, 'w', encoding='utf-8') as f:
            json.dump(json_records, f, indent=2, ensure_ascii=False)
        print(f"Saved JSON to: {out_path}")

    return json_records

if __name__ == '__main__':
    convert_excel_to_json()
